"""
ResMan Export Converter
========================
Flattens the raw grouped ResMan "Invoices by Location" and "Invoice Detail"
report exports (.xlsx, with report headers + Building/Unit group sections +
subtotal rows) into the flat CSVs that creekside_automation.py expects.

Usage:
    python convert_resman_exports.py \
        --location "Invoices_by_Location__June.xlsx" \
        --detail   "Invoice_Detail__June.xlsx" \
        --out-dir  /tmp/creekside_june
"""

import argparse
import csv
import os
import re
from datetime import datetime

import openpyxl

PROPERTY_NAME = "Oaks at Creekside"
GROUP_RE = re.compile(r"^(Building|Unit)\s*-\s*(\d+)$")


def convert_location(path, out_path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    current_unit = ""
    current_obj_type = ""
    out_rows = []

    for r in rows[6:]:  # skip 5 report-header lines + column header row
        invoice, acctg_date, gl, vendor, desc, qty, total = (list(r) + [None] * 7)[:7]

        if not isinstance(acctg_date, datetime):
            # Group header or subtotal row — not a transaction
            if isinstance(invoice, str):
                m = GROUP_RE.match(invoice.strip())
                if m:
                    kind, num = m.groups()
                    current_obj_type = kind
                    current_unit = num if kind == "Unit" else ""
                elif invoice.strip() == "None":
                    current_obj_type = ""
                    current_unit = ""
            continue

        out_rows.append({
            "PropertyName":  PROPERTY_NAME,
            "Unit#":         current_unit,
            "ObjectType":    current_obj_type or "Common Area",
            "Invoice #":     str(invoice or "").strip(),
            "Acctg Date":    acctg_date.strftime("%m/%d/%Y"),
            "GL Acc Number": str(gl or "").strip(),
            "Description":   str(desc or "").strip(),
            "Quantity":      qty if qty is not None else 1,
            "Total":         total if total is not None else 0,
            "Vendor":        str(vendor or "").strip(),
        })

    fieldnames = ["PropertyName", "Unit#", "ObjectType", "Invoice #",
                  "Acctg Date", "GL Acc Number", "Description", "Quantity", "Total", "Vendor"]
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)

    return out_rows


def convert_detail(path, out_path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    current_vendor = ""
    out_rows = []
    for r in rows[6:]:
        invoice, inv_date, _, acctg_date, due_date, desc, _, total = (list(r) + [None] * 8)[:8]

        if invoice is not None and not isinstance(inv_date, datetime):
            current_vendor = str(invoice).strip()
            continue

        if invoice is None or not isinstance(inv_date, datetime):
            continue  # GL-split sub-row

        out_rows.append({
            "InvoiceNumber":       str(invoice).strip(),
            "VendorName":          current_vendor,
            "VendorAbbreviation":  "",
            "InvoiceDate":         inv_date.strftime("%m/%d/%Y"),
            "AccountingDate":      acctg_date.strftime("%m/%d/%Y") if isinstance(acctg_date, datetime) else "",
            "DueDate":             due_date.strftime("%m/%d/%Y") if isinstance(due_date, datetime) else "",
            "Description":         str(desc or "").strip(),
            "Total":               total if total is not None else 0,
            "AmountPaid":          "",
        })

    fieldnames = ["InvoiceNumber", "VendorName", "VendorAbbreviation", "InvoiceDate",
                  "AccountingDate", "DueDate", "Description", "Total", "AmountPaid"]
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)

    return out_rows


def main():
    parser = argparse.ArgumentParser(description="Convert raw ResMan xlsx exports to flat CSVs")
    parser.add_argument("--location", required=True)
    parser.add_argument("--detail", required=True)
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args()

    os.makedirs(args.out_dir, exist_ok=True)
    loc_out = os.path.join(args.out_dir, "location.csv")
    det_out = os.path.join(args.out_dir, "detail.csv")

    loc_rows = convert_location(args.location, loc_out)
    det_rows = convert_detail(args.detail, det_out)

    print(f"Location: {len(loc_rows)} transaction rows -> {loc_out}")
    print(f"Detail:   {len(det_rows)} invoice rows -> {det_out}")


if __name__ == "__main__":
    main()
